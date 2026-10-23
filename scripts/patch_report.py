"""Patch an already-assembled Progress Report 2 docx in place.

Usage:  python scripts/patch_report.py [--doc reports/progress2/T2.A5.docx]

Two jobs:

1. Sync Parts C and D from their Markdown sources in docs/progress2/. Each replaces
   everything between its own heading and the next part's heading, so a section is
   regenerated whatever was there before -- the highlighted "PART C PENDING" placeholder
   build_progress_report.py writes when the source is missing, a section an earlier run
   inserted, or one baked in at assembly time from an older revision of the source.
2. Re-render the Part A paragraphs listed in CORRECTIONS, which were assembled from an
   earlier revision of docs/progress2/part_a_status.md.

This script exists only because the department template is not committed, so
`build_progress_report.py --template ...` cannot be re-run from a clean clone. Commit the
template and this script becomes unnecessary -- a full rebuild is the canonical route.

Re-running is safe: job 1 regenerates the same sections from the same sources, and job 2
skips a paragraph that already carries the new text.
"""
import argparse
import sys
from pathlib import Path

import docx

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "scripts"))

from build_progress_report import add_inline, insert_after, markdown_blocks  # noqa: E402

# Each section is replaced from its heading up to the next part's heading. The last
# section runs to the end of the document body.
SECTIONS = [
    ("Part C", "Part D", "part_c_work.md"),
    ("Part D", None, "part_d_risks.md"),
]

# (text identifying the stale paragraph, replacement Markdown). The replacements are the
# current wording in docs/progress2/part_a_status.md; keep the two in step.
CORRECTIONS = [
    (
        "runs/ contains two pipeline records and no model records",
        "**4.5 First Model Implementation, Complete 25%.** `models/generate_zeroshot.py` "
        "and `models/train_qlora.py` are merged and their CPU paths are tested, and "
        "`runs/` now holds four model records from a Tesla T4. **But no model has been "
        "trained on the real corpus and no scored run exists at three seeds**: those four "
        "records are a single-seed smoke test on synthetic data with a 0.5B stand-in "
        "model, and `docs/COMPUTE_BUDGET.md` states in terms that none of it is a "
        "reportable result. Issues 298-24 and 298-25 were previously marked Done without "
        "a result and were reset to In Progress on 22 October under 298-45. The scored run "
        "is 298-53, due 3 November, and the QLoRA fine-tune at three seeds is 298-54, due "
        "9 November. Both are blocked on GPU confirmation (298-51).",
    ),
    (
        "those six are being reset under 298-45",
        "A row was only marked above 80% on Complete if an artifact exists on `main` that "
        "a reader can open. Issues marked Done on the Linear board without a "
        "corresponding artifact were treated as incomplete. Eight such issues were found "
        "and reset on 22 October under 298-45 — 298-10, 298-15, 298-16, 298-17, "
        "298-21, 298-23, 298-24 and 298-25 — so that Part C, this table and the "
        "repository agree with each other. The audit is recorded in "
        "`docs/progress2/board_reset_298-45.md`.",
    ),
]


def find_heading(doc, text):
    for par in doc.paragraphs:
        if par.style.name.startswith("Heading") and par.text.strip().startswith(text):
            return par
    return None


def sync_section(doc, start_text, end_text, source: Path):
    start = find_heading(doc, start_text)
    if start is None:
        sys.exit(f"could not locate the {start_text!r} heading")

    body = start._p.getparent()
    children = list(body)
    lo = children.index(start._p)

    if end_text is None:
        # Last section: run to the end of the body, leaving the trailing sectPr in place.
        hi = len(children)
        while hi > lo + 1 and children[hi - 1].tag.endswith("}sectPr"):
            hi -= 1
    else:
        end = find_heading(doc, end_text)
        if end is None:
            sys.exit(f"could not locate the {end_text!r} heading")
        hi = children.index(end._p)
        if hi < lo:
            sys.exit(f"{end_text!r} appears before {start_text!r}")

    removed = children[lo + 1:hi]
    for el in removed:
        body.remove(el)

    insert_after(start._p, markdown_blocks(doc, source.read_text()))
    return f"{start_text} synced from {source.name} ({len(removed)} element(s) replaced)"


def rewrite(par, markdown):
    """Replace a paragraph's content, keeping the paragraph and its style."""
    for run in list(par.runs):
        run._r.getparent().remove(run._r)
    add_inline(par, markdown)


def patch(doc_path: Path, source_dir: Path):
    doc = docx.Document(doc_path)
    changed = [
        sync_section(doc, start, end, source_dir / name)
        for start, end, name in SECTIONS
    ]

    for needle, replacement in CORRECTIONS:
        par = next((p for p in doc.paragraphs if needle in p.text), None)
        if par is not None:
            rewrite(par, replacement)
            changed.append(f"corrected: {needle[:40]}...")

    doc.save(doc_path)
    for line in changed:
        print(f"  {line}")
    print(f"wrote {doc_path}")


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--doc", type=Path,
                    default=ROOT / "reports" / "progress2" / "T2.A5.docx")
    ap.add_argument("--source-dir", type=Path, default=ROOT / "docs" / "progress2",
                    help="directory holding the Part C and Part D Markdown sources")
    args = ap.parse_args()
    patch(args.doc, args.source_dir)
